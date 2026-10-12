"""Google es una proyección recuperable de SQLite, nunca la fuente de verdad.

Solo administra pestañas CRT_* con encabezados propios. No cambia permisos ni
borra pestañas del usuario. Las escrituras RAW evitan ejecutar fórmulas del portal.
"""
import hashlib
import json
import mimetypes
import os
import logging
from openpyxl import load_workbook
from .reporting import HEADERS, projection
from pathlib import Path
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google.oauth2.service_account import Credentials as ServiceCredentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaFileUpload
from .attachments import _digest
from .reporting import local_file_status

log = logging.getLogger(__name__)

SCOPES = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']

def atomic_token(path, text):
    path = Path(path)
    tmp = path.with_suffix('.tmp')
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, 'w') as f:
        f.write(text)
    os.chmod(tmp, 0o600)
    tmp.replace(path)

def credentials(path):
    if not path:
        raise ValueError('Configurar CRT_GOOGLE_CREDENTIALS con una ruta a JSON')
    data = json.loads(Path(path).read_text())
    if data.get('type') == 'service_account':
        return ServiceCredentials.from_service_account_info(data, scopes=SCOPES)
    creds = Credentials.from_authorized_user_info(data, scopes=SCOPES)
    if not creds.valid:
        if not creds.refresh_token:
            raise ValueError('Falta refresh_token; ejecutar scripts/authorize_google.py en estación de trabajo')
        creds.refresh(Request())
        atomic_token(path, creds.to_json())
    return creds

def execute(request):
    return request.execute(num_retries=5)

def validate_drive_folder(cfg, drive, creds):
    if not cfg.drive_folder_id:
        raise ValueError('Falta CRT_DRIVE_FOLDER_ID')
    try:
        folder = execute(drive.files().get(fileId=cfg.drive_folder_id,
                         fields='id,name,mimeType,trashed,driveId,capabilities(canAddChildren)', supportsAllDrives=True))
    except HttpError as exc:
        raise RuntimeError(f'Drive HTTP {exc.resp.status}: revisar ID de carpeta, permisos y Drive API') from exc
    if folder.get('mimeType') != 'application/vnd.google-apps.folder' or folder.get('trashed'):
        raise ValueError('El destino de Drive no es una carpeta válida')
    if isinstance(creds, ServiceCredentials) and not folder.get('driveId'):
        raise ValueError('Cuenta de servicio sin almacenamiento en Mi unidad: usar OAuth de usuario o una unidad compartida')
    if not folder.get('capabilities', {}).get('canAddChildren'):
        raise PermissionError('La credencial no puede agregar archivos a la carpeta de Drive')
    log.info('DRIVE | destino=%s | id=%s | acceso de escritura verificado', folder.get('name'), cfg.drive_folder_id)

def sync_drive(store, cfg, drive):
    if not cfg.drive_folder_id:
        raise ValueError('Falta CRT_DRIVE_FOLDER_ID')
    failures = []
    uploaded = reused = 0
    # Incluye archivados: un adjunto pendiente no desaparece al cerrar la consulta.
    for row in store.comments_for_export():
        for att in json.loads(row['adjuntos_json']):
            path = Path(att.get('local') or '')
            if local_file_status(att)[0] != 'Local verificado':
                log.error('DRIVE | folio=%03d consulta=%s archivo=%s | copia local no verificada', row['folio'], row['consulta_id'], att['nombre'])
                failures.append(att['nombre'])
                continue
            checksum = _digest(path)
            key = hashlib.sha256((cfg.drive_folder_id + '|' + att['url'] + '|' + checksum).encode()).hexdigest()
            saved = store.conn.execute('SELECT * FROM drive_files WHERE key=?', (key,)).fetchone()
            try:
                if saved and saved['url']:
                    reused += 1
                    continue
                if not saved:
                    file_id = execute(drive.files().generateIds(count=1, space='drive'))['ids'][0]
                    # Persistir ANTES del upload. Un reintento usa el mismo ID, incluso si se perdió la respuesta.
                    store.conn.execute('INSERT INTO drive_files VALUES (?,?,?,?)', (key, file_id, '', checksum))
                    store.conn.commit()
                else:
                    file_id = saved['file_id']
                try:
                    remote = execute(drive.files().get(fileId=file_id, fields='id,webViewLink', supportsAllDrives=True))
                except HttpError as exc:
                    if exc.resp.status != 404:
                        raise
                    remote = execute(drive.files().create(body={
                        'id': file_id, 'name': path.name, 'parents': [cfg.drive_folder_id],
                        'appProperties': {'crt_key': key, 'sha256': checksum}},
                        media_body=MediaFileUpload(str(path), mimetype=mimetypes.guess_type(path.name)[0] or 'application/octet-stream', resumable=True),
                        fields='id,webViewLink', supportsAllDrives=True))
                store.conn.execute('UPDATE drive_files SET url=? WHERE key=?',
                                   (remote.get('webViewLink') or f'https://drive.google.com/file/d/{file_id}/view', key))
                store.conn.commit()
                uploaded += 1
                log.info('DRIVE | folio=%03d consulta=%s archivo=%s | guardado | id=%s', row['folio'], row['consulta_id'], path.name, file_id)
            except Exception as exc:
                log.error('DRIVE | folio=%03d consulta=%s archivo=%s | FALLÓ | %s: %s', row['folio'], row['consulta_id'], path.name, type(exc).__name__, exc)
                failures.append(att['nombre'])
    log.info('DRIVE | nuevos=%s ya registrados=%s pendientes=%s', uploaded, reused, len(failures))
    if not failures:
        from .storage import now_iso
        store.set_meta('last_drive_synced_at', now_iso())
    return failures

def _cell(value):
    # SQLite conserva el texto íntegro. Sheets tiene límite de caracteres por celda.
    s = '' if value is None else str(value)
    return s if len(s) <= 48000 else s[:47900] + '\n[TRUNCADO; texto íntegro en SQLite]'

def workbook_rows(path):
    """Publicar exactamente las celdas del Excel generado, no reconstruir desde SQLite."""
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        for name, header in HEADERS.items():
            if name not in book.sheetnames:
                raise ValueError(f'Falta hoja local {name}')
            rows=[list(row) for row in book[name].iter_rows(values_only=True)]
            while len(rows)>1 and not any(v is not None for v in rows[-1]):rows.pop()
            if rows[0]!=header:
                raise ValueError(f'Encabezados locales incompatibles: {name}')
            yield 'CRT_'+name, [[v if isinstance(v,(int,float)) else _cell(v) for v in row] for row in rows]
    finally:
        book.close()

def sync_sheets(store, cfg, sheets):
    api = sheets.spreadsheets()
    book = execute(api.get(spreadsheetId=cfg.spreadsheet_id, fields='sheets(properties)'))
    existing = {s['properties']['title']: s['properties'] for s in book['sheets']}
    for name, rows in workbook_rows(cfg.xlsx_path):
        log.info("SHEETS | hoja=%s | filas=%s | origen=%s",name,len(rows)-1,cfg.xlsx_path)
        if name not in existing:
            result = execute(api.batchUpdate(spreadsheetId=cfg.spreadsheet_id, body={'requests': [
                {'addSheet': {'properties': {'title': name}}}]}))
            prop = result['replies'][0]['addSheet']['properties']
        else:
            prop = existing[name]
            current = execute(api.values().get(spreadsheetId=cfg.spreadsheet_id, range=f"'{name}'!1:1")).get('values', [])
            legacy_headers={'CRT_Comentarios':['ID comentario','ID consulta','Consulta','Autor','Fecha local México','Comentario','Adjuntos','Visible último barrido','Primera detección','Última vez visto'],
                            'CRT_Adjuntos':['ID comentario','ID consulta','Archivo','URL original','Ruta local','SHA-256','URL Drive','Estado']}
            if current and current[0] != rows[0] and current[0] != legacy_headers.get(name):
                raise ValueError(f'Encabezados incompatibles en {name}; no sobrescribir contenido ajeno')
        sid, count, width = prop['sheetId'], len(rows), len(rows[0])
        grid = prop.get('gridProperties', {})
        execute(api.batchUpdate(spreadsheetId=cfg.spreadsheet_id, body={'requests': [
            {'updateSheetProperties': {'properties': {'sheetId': sid, 'gridProperties': {
                'rowCount': max(count, grid.get('rowCount', 1000)), 'columnCount': max(width, grid.get('columnCount', 26)),
                'frozenRowCount': 1}}, 'fields': 'gridProperties'}}]}))
        # Bloques limitados por bytes. Reescribir posiciones fijas permite reanudar tras timeout sin duplicar.
        batch, size, start = [], 0, 1
        for row in rows:
            row_bytes = len(json.dumps(row, ensure_ascii=False).encode())
            if batch and size + row_bytes > 500000:
                execute(api.values().update(spreadsheetId=cfg.spreadsheet_id, range=f"'{name}'!A{start}", valueInputOption='RAW', body={'values': batch}))
                start += len(batch); batch, size = [], 0
            batch.append(row); size += row_bytes
        if batch:
            execute(api.values().update(spreadsheetId=cfg.spreadsheet_id, range=f"'{name}'!A{start}", valueInputOption='RAW', body={'values': batch}))
        # Limpiar solo filas sobrantes DESPUÉS de escribir; jamás limpiar toda la hoja primero.
        previous = int(store.get_meta('sheets_rows:' + cfg.spreadsheet_id + ':' + name, grid.get('rowCount', count)))
        if previous > count:
            execute(api.values().clear(spreadsheetId=cfg.spreadsheet_id, range=f"'{name}'!A{count+1}:Z{previous}", body={}))
        store.set_meta('sheets_rows:' + cfg.spreadsheet_id + ':' + name, count)
        requests = [
            {'repeatCell': {'range': {'sheetId': sid, 'startRowIndex': 0, 'endRowIndex': 1}, 'cell': {'userEnteredFormat': {
                'backgroundColor': {'red': .38, 'green': .07, 'blue': .20}, 'textFormat': {'bold': True, 'foregroundColor': {'red': 1, 'green': 1, 'blue': 1}},
                'wrapStrategy': 'WRAP'}}, 'fields': 'userEnteredFormat'}},
            {'repeatCell': {'range': {'sheetId': sid, 'startRowIndex': 1, 'endRowIndex': max(count, 2)}, 'cell': {'userEnteredFormat': {
                'verticalAlignment': 'TOP', 'wrapStrategy': 'WRAP'}}, 'fields': 'userEnteredFormat'}},
            {'updateDimensionProperties': {'range': {'sheetId': sid, 'dimension': 'COLUMNS', 'startIndex': 0, 'endIndex': width}, 'properties': {'pixelSize': 190}, 'fields': 'pixelSize'}},
            {'updateDimensionProperties': {'range': {'sheetId': sid, 'dimension': 'ROWS', 'startIndex': 0, 'endIndex': 1}, 'properties': {'pixelSize': 45}, 'fields': 'pixelSize'}},
            {'setBasicFilter': {'filter': {'range': {'sheetId': sid, 'startRowIndex': 0, 'endRowIndex': count, 'endColumnIndex': width}}}}
        ]
        if name == 'CRT_Comentarios':
            requests.append({'updateDimensionProperties': {'range': {'sheetId': sid, 'dimension': 'COLUMNS', 'startIndex': 5, 'endIndex': 6}, 'properties': {'pixelSize': 580}, 'fields': 'pixelSize'}})
        execute(api.batchUpdate(spreadsheetId=cfg.spreadsheet_id, body={'requests': requests}))

def synchronize(store, cfg):
    if not cfg.xlsx_path.is_file():
        raise FileNotFoundError('Primero generar Excel local')
    creds = credentials(cfg.google_credentials)
    failures=[]
    if cfg.sync_drive:
        log.info('DRIVE | Iniciando subida de adjuntos locales')
        drive=build('drive','v3',credentials=creds,cache_discovery=False)
        validate_drive_folder(cfg,drive,creds)
        failures=sync_drive(store,cfg,drive)
        # Incorporar los enlaces remotos ANTES de publicar el Excel en Sheets.
        from .excel import export_excel
        export_excel(store,cfg.xlsx_path,cfg)
    if failures:
        raise RuntimeError(f'{len(failures)} adjuntos pendientes de Drive')
    if cfg.sync_sheets:
        sync_sheets(store,cfg,build('sheets','v4',credentials=creds,cache_discovery=False))
        store.set_meta('last_synced_excel_sha256',_digest(cfg.xlsx_path))
        from .storage import now_iso
        store.set_meta('last_synced_at',now_iso())
        log.info('SHEETS | Sincronización completada | archivo=%s | sha256=%s',cfg.xlsx_path,_digest(cfg.xlsx_path))
