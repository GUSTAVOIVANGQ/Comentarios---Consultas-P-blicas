"""Prueba end-to-end: servidor local que imita el portal (contenido cargado por JS,
paginación de 10 en 10, adjuntos) + Chromium real. Se omite si Chromium no arranca."""
import json
import os
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest
from openpyxl import load_workbook

from crt_comments.config import Config
from crt_comments.runner import EXIT_OK, run_once
from crt_comments.storage import Store

PAGE_SIZE = 10
STATE = {}


def _make_comments(n):
    out = []
    for i in range(n, 0, -1):  # más nuevo primero, como el portal
        out.append({
            "autor": f"Persona {i}", "fecha": f"{(i % 28) + 1:02d}/10/2026 {(i % 23):02d}:{(i * 7) % 60:02d}",
            "texto": f"Comentario número {i}",
            "adjunto": f"{i:08d}-aaaa-bbbb-cccc-000000000000.pdf" if i % 5 == 0 else None,
        })
    return out


def _detail_html(cid, period_end):
    return f"""<html><body>
<h1>Consulta de prueba {cid}</h1><span>Consulta Activa</span>
<p><strong>Fecha de publicación:</strong> 10 de septiembre de 2026</p>
<p><strong>Periodo de la consulta:</strong> del 11 de septiembre de 2026, a las 12:00 a. m. al {period_end}</p>
<section><h2>Comentarios de la Ciudadanía</h2><span id="chip">0 comentarios</span>
<div id="list"><p>Cargando comentarios...</p></div>
<div><button id="prev">Anteriores</button><span id="pg">Página 1</span><button id="next">Siguientes</button></div></section>
<script>
const ID={cid}; let page=1;
async function load(p){{
  document.getElementById('list').innerHTML='<p>Cargando comentarios...</p>';
  const d=await (await fetch('/api/'+ID+'?page='+p)).json(); page=p;
  document.getElementById('chip').textContent=d.total+' comentarios';
  document.getElementById('pg').textContent='Página '+p+' de '+d.pages;
  document.getElementById('next').disabled = p>=d.pages;
  document.getElementById('prev').disabled = p<=1;
  document.getElementById('list').innerHTML = d.items.map(c =>
    '<div class="k"><div><b>'+c.autor+'</b> <span>'+c.fecha+'</span></div><p>'+c.texto+'</p>'+
    (c.adjunto ? '<div><div>'+c.adjunto+'</div><div>DOCUMENTO ADJUNTO</div><a href="/adjunto-comentario/'+c.adjunto+'">Descargar</a></div>' : '')+
    '</div>').join('');
}}
document.getElementById('next').onclick=()=>load(page+1);
document.getElementById('prev').onclick=()=>load(page-1);
load(1);
</script></body></html>"""


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="text/html; charset=utf-8"):
        data = body if isinstance(body, bytes) else body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        u = urlparse(self.path)
        parts = u.path.strip("/").split("/")
        if parts[:2] == ["ConsultaPublica", "Detalle"] and len(parts) == 3 and int(parts[2]) in STATE:
            cid = int(parts[2])
            return self._send(200, _detail_html(cid, STATE[cid]["period_end"]))
        if parts[0] == "api" and int(parts[1]) in STATE:
            cid = int(parts[1])
            page = int(parse_qs(u.query).get("page", ["1"])[0])
            items = STATE[cid]["comments"]
            chunk = items[(page - 1) * PAGE_SIZE: page * PAGE_SIZE]
            pages = max(1, -(-len(items) // PAGE_SIZE))
            return self._send(200, json.dumps({"total": len(items), "pages": pages, "items": chunk}),
                              "application/json")
        if parts[0] == "adjunto-comentario":
            return self._send(200, b"%PDF-1.4 fake", "application/pdf")
        return self._send(404, "<html><body><h1>No encontrado</h1></body></html>")


@pytest.fixture()
def server():
    STATE.clear()
    STATE[24] = {"period_end": "9 de octubre de 2099, a las 11:59 p. m.", "comments": _make_comments(32)}
    STATE[25] = {"period_end": "1 de enero de 2020, a las 11:59 p. m.", "comments": _make_comments(3)}
    STATE[26] = {"period_end": "9 de octubre de 2099, a las 11:59 p. m.", "comments": []}
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()


def _cfg(tmp_path, base):
    return Config(base_url=base, start_id=24, max_consecutive_misses=3, data_dir=tmp_path / "data",
                  log_dir=tmp_path / "logs", page_delay_s=0.01, load_timeout_s=15, discover_index=False, probe_batch=4, browser_sandbox=False, browser_executable=os.environ.get("CRT_TEST_BROWSER", ""))


def test_full_flow(server, tmp_path):
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            p.chromium.launch(executable_path=os.environ.get("CRT_TEST_BROWSER") or None).close()
    except Exception as exc:  # noqa: BLE001
        pytest.skip(f"Chromium no disponible: {exc}")

    cfg = _cfg(tmp_path, server)

    # ---- corrida 1: descubre 24, 25, 26; 27/28/29 no existen -> se detiene ----
    assert run_once(cfg) == EXIT_OK
    st = Store(cfg.db_path)
    k = st.known_consultas()
    assert sorted(k) == [24, 25, 26]
    assert st.count_comments(24) == 32 and k[24]["total_reportado"] == 32 and k[24]["validacion_ok"] == 1
    assert st.count_comments(25) == 3
    assert st.count_comments(26) == 0 and k[26]["validacion_ok"] == 1     # consulta vacía válida
    assert k[24]["finalizada"] == 0          # periodo abierto
    assert k[25]["finalizada"] == 1          # periodo cerrado y validado
    # adjuntos descargados (i%5==0 -> 5,10,...,30 = 6 en la consulta 24)
    assert len(list((cfg.attachments_dir / "24").glob("*.pdf"))) == 6
    st.close()

    wb = load_workbook(cfg.xlsx_path)
    assert wb["Comentarios"].max_row == 1 + 32 + 3
    assert wb["Consultas"]["L2"].value == "OK"

    # ---- corrida 2: llega un comentario nuevo en la 24 -----------------------
    STATE[24]["comments"] = _make_comments(33)
    STATE[25]["comments"] = _make_comments(4)
    assert run_once(cfg) == EXIT_OK
    st = Store(cfg.db_path)
    assert st.count_comments(24) == 33                                    # +1, sin duplicar
    assert st.conn.execute("SELECT comentarios_nuevos FROM ejecuciones ORDER BY id DESC LIMIT 1").fetchone()[0] == 2
    # La 25 cerrada se revisa: captura publicación tardía.
    assert st.count_comments(25) == 4
    st.close()
    assert load_workbook(cfg.xlsx_path)["Comentarios"].max_row == 1 + 33 + 4

    # ---- corrida 3: aparece la consulta 27 (id nuevo) ------------------------
    STATE[27] = {"period_end": "9 de octubre de 2099, a las 11:59 p. m.", "comments": _make_comments(11)}
    assert run_once(cfg, only_ids=[24, 25, 26, 27]) == EXIT_OK
    st = Store(cfg.db_path)
    assert st.count_comments(27) == 11        # 2 páginas (10 + 1)
    st.close()
