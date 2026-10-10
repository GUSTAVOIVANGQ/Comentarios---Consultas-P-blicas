from openpyxl import load_workbook

from crt_comments.excel import export_excel
from crt_comments.models import assign_keys
from crt_comments.parsers import parse_detail_html
from crt_comments.storage import Store
from pathlib import Path

FX = Path(__file__).parent / "fixtures"


def test_upsert_is_idempotent_and_excel_roundtrip(tmp_path):
    snap = parse_detail_html((FX / "detalle_pagina2.html").read_text(encoding="utf-8"))
    assign_keys(24, snap.comments)
    st = Store(tmp_path / "t.sqlite3")
    assert st.upsert_comments(24, snap.comments) == 4
    assert st.upsert_comments(24, snap.comments) == 0          # segunda corrida: nada nuevo
    st.upsert_consulta(24, "https://x/24", snap, True, False, "")
    run = st.start_run(); st.finish_run(run, 1, 4, 0, 0, "ok")

    out = tmp_path / "out.xlsx"
    export_excel(st, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Resumen", "Comentarios", "Adjuntos", "Consultas", "Ejecuciones"]
    ws = wb["Comentarios"]
    assert ws.max_row == 5                                      # encabezado + 4
    assert ws["D2"].value == "Luis Alberto Arce Gómez"
    assert wb["Adjuntos"]["D2"].hyperlink is not None                       # adjunto único -> hipervínculo
    c = wb["Consultas"]
    assert c["G2"].value == 32 and c["H2"].value == 4 and c["L2"].value == "REVISAR"  # 4 != 32
