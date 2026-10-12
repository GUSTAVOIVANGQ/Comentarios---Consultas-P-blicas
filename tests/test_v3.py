"""Pruebas del flujo local, folios permanentes y límites del panel."""
import json
import sqlite3
import threading
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import Mock
import requests
import pytest
from openpyxl import load_workbook
from crt_comments.config import Config
from crt_comments.storage import Store,SCHEMA
from crt_comments.models import Comment,Attachment,DetailSnapshot,assign_keys
from crt_comments.scraper import ScrapeResult
from crt_comments.excel import export_excel
from crt_comments.runner import publish_cloud,run_once
from crt_comments.bundle import create_bundle
from crt_comments.ui import make_server
from crt_comments.discovery import enable_closed


def c(text='Comentario',atts=None):return Comment('Autor','09/10/2026 14:29',None,text,atts or [])

def cfg(tmp):return Config(data_dir=tmp/'data',log_dir=tmp/'logs',page_delay_s=0,download_attachments=False,sync_drive=False)

def save(st,comments):
    assign_keys(24,comments)
    snap=DetailSnapshot('Prueba',total_reportado=len(comments),page_current=1,page_total=1)
    return st.save_validated(24,'https://portal.crt.gob.mx/ConsultaPublica/Detalle/24',ScrapeResult(24,True,snap,comments,1,1,True),False)


def test_folios_never_renumber_or_skip_on_unchanged_run(tmp_path):
    st=Store(tmp_path/'db')
    save(st,[c('A'),c('B')]);save(st,[c('B'),c('A')]);save(st,[c('B'),c('A'),c('C')])
    assert [(r['comentario'],r['folio']) for r in st.comments_for_export()]==[('A',1),('B',2),('C',3)]
    save(st,[c('C')]);save(st,[c('C'),c('D')])
    assert st.comments_for_export()[-1]['folio']==4
    st.close()


def test_migrate_original_database(tmp_path):
    db=tmp_path/'old.sqlite'
    conn=sqlite3.connect(db);conn.executescript(SCHEMA)
    conn.execute("INSERT INTO comentarios(hash,consulta_id,autor,comentario,primera_deteccion) VALUES ('original',24,'Autor','Texto','2026-10-09')")
    conn.commit();conn.close()
    st=Store(db);assert st.comments_for_export()[0]['folio']==1;st.close()
    st=Store(db);assert st.comments_for_export()[0]['folio']==1;st.close()


def test_hidden_checkbox_uses_native_click_and_verifies(tmp_path):
    toggle=Mock();toggle.count.return_value=1;toggle.is_checked.side_effect=[False,True];toggle.is_visible.return_value=False
    label=Mock();label.count.return_value=0
    page=Mock();page.locator.side_effect=[toggle,label]
    enable_closed(page)
    toggle.evaluate.assert_called_once_with('(element) => element.click()')
    toggle.check.assert_not_called()


def test_pending_file_blocks_cloud_before_auth(tmp_path,monkeypatch):
    import crt_comments.google_sync as google
    st=Store(tmp_path/'db');save(st,[c(atts=[Attachment('a.pdf','https://portal.crt.gob.mx/adjunto-comentario/a.pdf')])])
    sync=Mock();monkeypatch.setattr(google,'synchronize',sync)
    with pytest.raises(RuntimeError,match='001'):publish_cloud(st,cfg(tmp_path))
    sync.assert_not_called();st.close()


def test_download_failure_still_creates_local_excel(tmp_path,monkeypatch):
    import crt_comments.runner as runner
    conf=cfg(tmp_path);conf.download_attachments=True;conf.sync_sheets=True
    item=c(atts=[Attachment('a.pdf','https://portal.crt.gob.mx/adjunto-comentario/a.pdf')]);assign_keys(24,[item])
    snap=DetailSnapshot('Prueba',total_reportado=1,page_total=1,page_current=1)
    monkeypatch.setattr(runner,'BrowserSession',lambda cfg:nullcontext(Mock()))
    monkeypatch.setattr(runner,'_scrape_with_retries',lambda *args:ScrapeResult(24,True,snap,[item],1,1,True))
    monkeypatch.setattr(runner,'download_all',lambda *args:1)
    assert run_once(conf,only_ids=[24])==2
    book=load_workbook(conf.xlsx_path)
    assert book['Comentarios']['A2'].value=='001'
    assert book['Comentarios']['H2'].value=='Pendiente (0/1)'
    book.close()


def test_bundle_preserves_relative_excel_links(tmp_path):
    import hashlib,zipfile
    conf=cfg(tmp_path);conf.ensure_dirs();st=Store(conf.db_path)
    file=conf.attachments_dir/'24'/'001_a.pdf';file.parent.mkdir(parents=True);file.write_bytes(b'%PDF-test')
    file.with_suffix('.pdf.sha256').write_text(hashlib.sha256(file.read_bytes()).hexdigest())
    save(st,[c(atts=[Attachment('a.pdf','https://portal.crt.gob.mx/adjunto-comentario/a.pdf',str(file))])]);st.close()
    path=create_bundle(conf)
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        assert 'adjuntos/24/001_a.pdf' in z.namelist()
        manifest=json.loads(z.read('manifest.json'))
        assert manifest['adjuntos_incluidos']==1 and not manifest['pendientes']
    book=load_workbook(conf.xlsx_path)
    assert book['Adjuntos']['E2'].hyperlink.target=='adjuntos/24/001_a.pdf'
    assert book['Comentarios']['A2'].value=='001'
    book.close()


@pytest.fixture
def panel(tmp_path):
    conf=cfg(tmp_path);conf.ensure_dirs();st=Store(conf.db_path)
    save(st,[c('<script>alert(1)</script>')]);export_excel(st,conf.xlsx_path,conf);st.close()
    server=make_server(conf,port=0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    yield conf,server,f'http://127.0.0.1:{server.server_port}'
    server.shutdown();server.server_close()


def test_ui_read_download_and_csrf(panel):
    conf,server,url=panel
    assert requests.get(url+'/api/state').status_code==401
    s=requests.Session();assert s.get(url).status_code==200
    state=s.get(url+'/api/state').json();assert state['comments']==1 and state['excel_ready']
    assert s.get(url+'/api/comments').json()['rows'][0]['folio_label']=='001'
    assert s.get(url+'/download/excel').content[:2]==b'PK'
    assert s.post(url+'/api/jobs',json={'action':'export'}).status_code==403
    token=s.get(url+'/api/token').json()['token']
    assert s.post(url+'/api/jobs',headers={'X-CSRF-Token':token},json={'action':'bad'}).status_code==400
    assert s.get(url+'/api/state',headers={'Host':'attacker.test'}).status_code==403


def test_ui_rejects_attachment_path_traversal(panel,tmp_path):
    conf,server,url=panel
    st=Store(conf.db_path);st.conn.execute('UPDATE comentarios SET adjuntos_json=?',(json.dumps([{'nombre':'outside','url':'https://x','local':str(tmp_path/'private.txt')}]),));st.conn.commit();st.close()
    s=requests.Session();s.get(url)
    assert s.get(url+'/download/attachment/1/0').status_code==403


def test_sync_reads_excel_content_and_keeps_folio_text(tmp_path):
    from crt_comments.google_sync import workbook_rows
    conf=cfg(tmp_path);st=Store(conf.db_path);save(st,[c('Texto del Excel')]);export_excel(st,conf.xlsx_path,conf)
    save(st,[c('Texto del Excel'),c('Aún no exportado')])
    rows=dict(workbook_rows(conf.xlsx_path))['CRT_Comentarios']
    assert len(rows)==2 and rows[1][0]=='001' and rows[1][5]=='Texto del Excel'
    st.close()


def test_ui_starts_real_local_export_worker(panel):
    import time
    conf,server,url=panel
    s=requests.Session();s.get(url)
    token=s.get(url+'/api/token').json()['token']
    response=s.post(url+'/api/jobs',headers={'X-CSRF-Token':token},json={'action':'export'})
    assert response.status_code==202
    deadline=time.monotonic()+10
    while time.monotonic()<deadline:
        status=s.get(url+'/api/state').json()
        if status['jobs'][0]['status']!='En ejecución':break
        time.sleep(.05)
    assert status['jobs'][0]['status']=='Completada'
    assert 'EXCEL | guardado=' in s.get(url+'/api/logs').json()['text']
    book=load_workbook(conf.xlsx_path)
    assert book['Comentarios']['A2'].value=='001'
    book.close()
